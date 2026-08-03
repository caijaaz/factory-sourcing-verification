# -*- coding: utf-8 -*-
"""
工厂资质汇总表生成模板（openpyxl）。
用法：替换 ROWS / TRADERS 数据后运行：
  python gen_factory_xlsx.py [输出路径.xlsx]
表1「真实生产厂池」9 列；表2「剔除名单与方法」。
样式：表头深蓝底白字、冻结首行、自动筛选、wrap+top 对齐、细边框。
列结构（9 列，与品类无关，通用）：
  厂家 / 存续 / 进出口资质 / 生产资质 / 目标产品在产 / 定位 / 联系方式 / 地址 / 生产规模
"""
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUT = sys.argv[1] if len(sys.argv) > 1 else "工厂资质汇总.xlsx"

HEADERS = ["厂家", "存续", "进出口资质", "生产资质", "目标产品在产", "定位",
           "联系方式", "地址", "生产规模"]

# ===== 替换为实际核验数据（每行 9 列，与 HEADERS 对应）=====
ROWS = [
    ["示例工厂有限公司",
     "存续（在营）",
     "海关注册编码 XXXXXXXXXX ✓",
     "✓ 许可项目：{品类}生产；许可证 XX 号，有效期至 YYYY-MM-DD",
     "✓ 公开产品清单明确在产 / 待电话确认",
     "★ 首选 / Tier-1 / Tier-2 / Tier-3",
     "电话；邮箱；官网",
     "省市区详细地址",
     "产能、产线、主营品种"],
]

# ===== 剔除名单：[厂家, 剔除原因] =====
TRADERS = [
    ["示例贸易公司", "经营范围含'不含{品类}'，仅经营许可，无生产许可——贸易商"],
]

METHODS = [
    ["WebSearch 多引擎交叉", "工商经营范围 + 官网声明 + 许可证/海关注册编码，落实存续/进出口/生产资质/在产四维度"],
    ["政府/监管《生产企业名录》", "通过许可证编号、许可范围确认生产资质"],
    ["产业园区入驻名单", "确认属地与资质背书"],
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "真实生产厂池"

hdr_fill = PatternFill("solid", fgColor="1F4E78")
hdr_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")

ws.append(HEADERS)
for c in range(1, len(HEADERS) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = border

for r in ROWS:
    ws.append(r)

for ri in range(2, len(ROWS) + 2):
    for ci in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=ri, column=ci)
        cell.alignment = wrap_top
        cell.border = border
        if ci == 1:
            cell.font = Font(bold=True, size=11)

for i, w in enumerate([22, 12, 30, 22, 36, 26, 34, 34, 38], start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(HEADERS))}{len(ROWS)+1}"
ws.row_dimensions[1].height = 30

ws2 = wb.create_sheet("剔除名单与方法")
ws2.append(["一、剔除的贸易型线索（不满足硬性资质门槛）"])
ws2.append(["厂家", "剔除原因"])
for t in TRADERS:
    ws2.append(t)
ws2.append([])
ws2.append(["二、验证方法"])
ws2.append(["渠道", "用途"])
for m in METHODS:
    ws2.append(m)

ws2.cell(row=1, column=1).font = Font(bold=True, size=12, color="1F4E78")
for rr in range(2, 2 + len(TRADERS) + 1):
    for cc in (1, 2):
        ws2.cell(row=rr, column=cc).alignment = wrap_top
        ws2.cell(row=rr, column=cc).border = border
sec2 = 2 + len(TRADERS) + 2
ws2.cell(row=sec2, column=1).font = Font(bold=True, size=12, color="1F4E78")
for rr in range(sec2 + 1, sec2 + 2 + len(METHODS)):
    for cc in (1, 2):
        ws2.cell(row=rr, column=cc).alignment = wrap_top
        ws2.cell(row=rr, column=cc).border = border
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 80

wb.save(OUT)
print("saved:", OUT)
